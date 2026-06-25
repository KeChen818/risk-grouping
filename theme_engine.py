from __future__ import annotations

import json
import os
import re
from collections import Counter
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, IO

import numpy as np
import pandas as pd
import yaml
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity


REQUIRED_COLUMNS = [
    "risk_id",
    "risk_title",
    "risk_description",
    "business_unit",
    "risk_type",
]

OPTIONAL_COLUMNS = [
    "taxonomy",
    "root_cause",
    "likelihood",
    "impact",
    "materiality",
    "existing_rating",
]

INPUT_COLUMNS = REQUIRED_COLUMNS + OPTIONAL_COLUMNS

OUTPUT_COLUMNS = [
    "risk_id",
    "risk_title",
    "business_unit",
    "risk_type",
    "taxonomy",
    "root_cause",
    "primary_theme_id",
    "primary_theme",
    "primary_theme_score",
    "secondary_theme",
    "secondary_theme_score",
    "score_gap",
    "review_flag",
    "review_reason",
    "assignment_rationale",
    "llm_assignment_justification",
    "llm_evidence",
    "human_override_theme",
    "final_theme",
]

SUMMARY_COLUMNS = [
    "theme_name",
    "domain",
    "total_risks",
    "material_risks",
    "business_units",
    "risk_types",
    "top_root_causes",
    "key_risk_titles",
    "ai_summary",
    "llm_summary",
    "llm_key_drivers",
    "llm_management_takeaway",
]

LIST_FIELDS = [
    "typical_business_units",
    "typical_risk_types",
    "root_causes",
    "keywords",
]

COLUMN_VARIANTS = {
    "risk_id": {"risk id", "risk identifier", "risk number", "risk no", "id", "risk ref"},
    "risk_title": {"risk title", "risk name", "risk event", "title", "name"},
    "risk_description": {
        "risk description",
        "description",
        "risk statement",
        "risk detail",
        "event description",
        "details",
    },
    "business_unit": {
        "business unit",
        "business",
        "function",
        "division",
        "unit",
        "owner business",
        "owning business",
    },
    "risk_type": {"risk type", "type", "risk category", "category", "primary risk type"},
    "taxonomy": {
        "taxonomy",
        "risk taxonomy",
        "sub taxonomy",
        "sub-taxonomy",
        "risk subcategory",
        "sub category",
    },
    "root_cause": {"root cause", "driver", "cause", "risk driver", "underlying cause"},
    "likelihood": {"likelihood", "probability", "frequency"},
    "impact": {"impact", "severity", "consequence"},
    "materiality": {"materiality", "material", "material risk", "is material"},
    "existing_rating": {"existing rating", "rating", "risk rating", "score", "existing score"},
}

TOKEN_PATTERN = re.compile(r"[a-z0-9]+")


@dataclass
class VectorStoreStatus:
    backend: str
    index_path: str
    metadata_path: str
    theme_count: int
    ready: bool
    message: str


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass
    return " ".join(str(value).replace("\n", " ").split())


def _normalize_text(value: Any) -> str:
    return _clean_text(value).lower().replace("&", " and ")


def _tokens(value: Any) -> set[str]:
    return set(TOKEN_PATTERN.findall(_normalize_text(value)))


def _as_list(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, list):
        return [_clean_text(item) for item in value if _clean_text(item)]
    if isinstance(value, str):
        if not value.strip():
            return []
        return [item.strip() for item in value.split(",") if item.strip()]
    return [_clean_text(value)] if _clean_text(value) else []


def _join_values(values: list[Any]) -> str:
    return " ".join(_clean_text(value) for value in values if _clean_text(value))


def _canonicalize_column(column_name: str) -> str:
    normalized = (
        str(column_name)
        .strip()
        .lower()
        .replace("\n", " ")
        .replace("-", " ")
        .replace("/", " ")
        .replace(".", " ")
    )
    normalized = " ".join(normalized.split())
    snake = normalized.replace(" ", "_")

    for canonical, variants in COLUMN_VARIANTS.items():
        if snake == canonical or normalized in variants or snake in variants:
            return canonical
    return snake


class ThemeEngine:
    """Demo-friendly theme assignment engine with FAISS retrieval and TF-IDF fallback."""

    def __init__(
        self,
        base_dir: str | Path | None = None,
        embedding_model_name: str = "all-MiniLM-L6-v2",
    ) -> None:
        self.base_dir = Path(base_dir or Path(__file__).resolve().parent)
        self.theme_library_path = self.base_dir / "data" / "risk_theme_library.yaml"
        self.demo_inventory_path = self.base_dir / "data" / "demo_inventory.xlsx"
        self.vector_store_dir = self.base_dir / "vector_store"
        self.index_path = self.vector_store_dir / "theme_index.faiss"
        self.metadata_path = self.vector_store_dir / "theme_metadata.json"
        self.embedding_model_name = embedding_model_name

        self.themes_df = pd.DataFrame()
        self.theme_metadata: list[dict[str, Any]] = []
        self.theme_documents: list[str] = []
        self.theme_name_to_domain: dict[str, str] = {}

        self.faiss = None
        self.embedding_model = None
        self.faiss_index = None
        self.vectorizer: TfidfVectorizer | None = None
        self.theme_tfidf_matrix = None

        self.backend = "Not initialized"
        self.ready = False
        self.status_message = ""

    def initialize(self, rebuild: bool = False) -> VectorStoreStatus:
        self.themes_df = self.load_theme_library()
        self.theme_metadata = self._theme_metadata_from_library(self.themes_df)
        self.theme_documents = [item["theme_document"] for item in self.theme_metadata]
        self.theme_name_to_domain = {
            item["theme_name"]: item["domain"] for item in self.theme_metadata
        }

        if not rebuild and self.index_path.exists() and self.metadata_path.exists():
            if self._load_faiss_vector_store():
                return self.vector_store_status()

        if self._build_faiss_vector_store():
            return self.vector_store_status()

        self._build_tfidf_fallback()
        return self.vector_store_status()

    def rebuild_vector_store(self) -> VectorStoreStatus:
        return self.initialize(rebuild=True)

    def vector_store_status(self) -> VectorStoreStatus:
        return VectorStoreStatus(
            backend=self.backend,
            index_path=str(self.index_path),
            metadata_path=str(self.metadata_path),
            theme_count=len(self.theme_metadata),
            ready=self.ready,
            message=self.status_message,
        )

    def load_inventory(self, file_or_path: str | Path | IO[bytes]) -> pd.DataFrame:
        name = getattr(file_or_path, "name", str(file_or_path))
        suffix = Path(name).suffix.lower()
        if suffix in {".xlsx", ".xls"}:
            raw = pd.read_excel(file_or_path)
        else:
            raw = pd.read_csv(file_or_path)
        return self.normalize_inventory(raw)

    def load_demo_inventory(self) -> pd.DataFrame:
        return self.load_inventory(self.demo_inventory_path)

    def normalize_inventory(self, raw: pd.DataFrame) -> pd.DataFrame:
        df = raw.copy()
        rename_map = {}
        seen = set()
        for column in df.columns:
            canonical = _canonicalize_column(column)
            if canonical in INPUT_COLUMNS and canonical not in seen:
                rename_map[column] = canonical
                seen.add(canonical)
        df = df.rename(columns=rename_map)

        missing = [column for column in REQUIRED_COLUMNS if column not in df.columns]
        if missing:
            raise ValueError(f"Missing required input columns: {', '.join(missing)}")

        for column in OPTIONAL_COLUMNS:
            if column not in df.columns:
                df[column] = ""
        for column in INPUT_COLUMNS:
            df[column] = df[column].fillna("").astype(str).str.strip()

        extra_columns = [column for column in df.columns if column not in INPUT_COLUMNS]
        return df[INPUT_COLUMNS + extra_columns]

    def load_theme_library(self) -> pd.DataFrame:
        with self.theme_library_path.open("r", encoding="utf-8") as file:
            records = yaml.safe_load(file) or []
        if not isinstance(records, list):
            raise ValueError("risk_theme_library.yaml must contain a YAML list of theme cards.")

        themes = pd.DataFrame(records)
        required = ["theme_id", "theme_name", "domain", "description"]
        missing = [column for column in required if column not in themes.columns]
        if missing:
            raise ValueError(f"Theme library missing fields: {', '.join(missing)}")

        for column in required:
            themes[column] = themes[column].fillna("").astype(str).str.strip()
        for column in LIST_FIELDS:
            if column not in themes.columns:
                themes[column] = [[] for _ in range(len(themes))]
            themes[column] = themes[column].apply(_as_list)

        return themes

    def theme_library_for_display(self) -> pd.DataFrame:
        display = self.themes_df.copy()
        if display.empty:
            display = self.load_theme_library()
        for column in LIST_FIELDS:
            display[column] = display[column].apply(lambda values: ", ".join(_as_list(values)))
        return display

    def build_theme_document(self, theme: pd.Series | dict[str, Any]) -> str:
        return _join_values(
            [
                theme.get("theme_id", ""),
                theme.get("theme_name", ""),
                theme.get("domain", ""),
                theme.get("description", ""),
                " ".join(_as_list(theme.get("typical_business_units", []))),
                " ".join(_as_list(theme.get("typical_risk_types", []))),
                " ".join(_as_list(theme.get("root_causes", []))),
                " ".join(_as_list(theme.get("keywords", []))),
            ]
        )

    def build_risk_profile_text(self, risk: pd.Series | dict[str, Any]) -> str:
        return _join_values(
            [
                risk.get("risk_title", ""),
                risk.get("risk_description", ""),
                risk.get("business_unit", ""),
                risk.get("risk_type", ""),
                risk.get("taxonomy", ""),
                risk.get("root_cause", ""),
            ]
        )

    def assign_inventory(
        self,
        inventory: pd.DataFrame,
        auto_threshold: float = 0.80,
        review_threshold: float = 0.70,
        ambiguity_gap: float = 0.10,
    ) -> tuple[pd.DataFrame, pd.DataFrame]:
        if not self.ready:
            self.initialize()

        risks = self.normalize_inventory(inventory)
        assigned_records = []
        candidate_records = []

        for _, risk in risks.iterrows():
            risk_profile = self.build_risk_profile_text(risk)
            candidates = self.search_theme_candidates(risk_profile, top_k=3)
            scored_candidates = [
                self._score_candidate(risk, risk_profile, candidate) for candidate in candidates
            ]
            scored_candidates = sorted(
                scored_candidates,
                key=lambda item: item["final_score"],
                reverse=True,
            )

            top = scored_candidates[0]
            second = scored_candidates[1] if len(scored_candidates) > 1 else top
            top_score = float(top["final_score"])
            second_score = float(second["final_score"])
            score_gap = top_score - second_score

            if top_score >= auto_threshold and score_gap >= ambiguity_gap:
                review_flag = "No"
                review_reason = ""
                secondary_theme = ""
                secondary_theme_score = None
            elif top_score >= review_threshold and score_gap < ambiguity_gap:
                review_flag = "Yes"
                review_reason = "Multiple plausible themes"
                secondary_theme = second["theme_name"]
                secondary_theme_score = round(second_score, 3)
            else:
                review_flag = "Yes"
                review_reason = "Low confidence"
                secondary_theme = ""
                secondary_theme_score = None

            assigned_records.append(
                {
                    "risk_id": risk["risk_id"],
                    "risk_title": risk["risk_title"],
                    "risk_description": risk["risk_description"],
                    "business_unit": risk["business_unit"],
                    "risk_type": risk["risk_type"],
                    "taxonomy": risk["taxonomy"],
                    "root_cause": risk["root_cause"],
                    "materiality": risk.get("materiality", ""),
                    "primary_theme_id": top["theme_id"],
                    "primary_theme": top["theme_name"],
                    "primary_theme_domain": top["domain"],
                    "primary_theme_score": round(top_score, 3),
                    "secondary_theme": secondary_theme,
                    "secondary_theme_score": secondary_theme_score,
                    "score_gap": round(score_gap, 3),
                    "review_flag": review_flag,
                    "review_reason": review_reason,
                    "assignment_rationale": self._assignment_rationale(top),
                    "llm_assignment_justification": "",
                    "llm_evidence": "",
                    "human_override_theme": "",
                    "final_theme": top["theme_name"],
                }
            )

            for rank, candidate in enumerate(scored_candidates, start=1):
                candidate_records.append(
                    {
                        "risk_id": risk["risk_id"],
                        "candidate_rank": rank,
                        "theme_id": candidate["theme_id"],
                        "theme_name": candidate["theme_name"],
                        "domain": candidate["domain"],
                        "semantic_similarity": round(candidate["semantic_similarity"], 3),
                        "business_match_score": round(candidate["business_match_score"], 3),
                        "risk_type_match_score": round(candidate["risk_type_match_score"], 3),
                        "root_cause_match_score": round(candidate["root_cause_match_score"], 3),
                        "keyword_match_score": round(candidate["keyword_match_score"], 3),
                        "final_score": round(candidate["final_score"], 3),
                    }
                )

        assigned = pd.DataFrame(assigned_records)
        candidates = pd.DataFrame(candidate_records)
        return assigned, candidates

    def search_theme_candidates(self, risk_profile: str, top_k: int = 3) -> list[dict[str, Any]]:
        if self.faiss_index is not None and self.embedding_model is not None:
            vector = self.embedding_model.encode(
                [risk_profile],
                normalize_embeddings=True,
                show_progress_bar=False,
            ).astype("float32")
            scores, indexes = self.faiss_index.search(vector, min(top_k, len(self.theme_metadata)))
            candidates = []
            for raw_score, index in zip(scores[0], indexes[0]):
                if index < 0:
                    continue
                theme = dict(self.theme_metadata[int(index)])
                theme["semantic_similarity"] = float(np.clip((raw_score + 1.0) / 2.0, 0.0, 1.0))
                candidates.append(theme)
            return candidates

        if self.vectorizer is None or self.theme_tfidf_matrix is None:
            self._build_tfidf_fallback()

        risk_vector = self.vectorizer.transform([risk_profile])
        raw_scores = cosine_similarity(risk_vector, self.theme_tfidf_matrix)[0]
        calibrated = np.sqrt(np.clip(raw_scores, 0.0, 1.0)) * 1.15
        top_indexes = np.argsort(calibrated)[::-1][:top_k]

        candidates = []
        for index in top_indexes:
            theme = dict(self.theme_metadata[int(index)])
            theme["semantic_similarity"] = float(np.clip(calibrated[index], 0.0, 1.0))
            candidates.append(theme)
        return candidates

    def generate_theme_summary(self, assigned: pd.DataFrame) -> pd.DataFrame:
        if assigned.empty:
            return pd.DataFrame(columns=SUMMARY_COLUMNS)

        working = assigned.copy()
        if "final_theme" not in working.columns:
            working["final_theme"] = working["primary_theme"]

        rows = []
        for theme_name, group in working.groupby("final_theme", dropna=False):
            domain = self.theme_name_to_domain.get(
                theme_name,
                group.get("primary_theme_domain", pd.Series([""])).iloc[0],
            )
            business_units = self._join_unique(group["business_unit"])
            risk_types = self._join_unique(group["risk_type"])
            root_causes = self._top_values(group["root_cause"], limit=4)
            material_risks = int(group.get("materiality", pd.Series(dtype=str)).apply(self._is_material).sum())
            key_titles = "; ".join(_clean_text(value) for value in group["risk_title"].head(5))

            rows.append(
                {
                    "theme_name": theme_name,
                    "domain": domain,
                    "total_risks": int(len(group)),
                    "material_risks": material_risks,
                    "business_units": business_units,
                    "risk_types": risk_types,
                    "top_root_causes": root_causes,
                    "key_risk_titles": key_titles,
                    "ai_summary": self._rule_based_summary(
                        theme_name,
                        group,
                        business_units,
                        risk_types,
                        root_causes,
                        material_risks,
                    ),
                    "llm_summary": "",
                    "llm_key_drivers": "",
                    "llm_management_takeaway": "",
                }
            )

        return pd.DataFrame(rows)[SUMMARY_COLUMNS].sort_values(
            ["material_risks", "total_risks", "theme_name"],
            ascending=[False, False, True],
        )

    def build_excel_export(
        self,
        assigned: pd.DataFrame,
        review_queue: pd.DataFrame,
        theme_summary: pd.DataFrame,
    ) -> bytes:
        output = BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            assigned[OUTPUT_COLUMNS].to_excel(writer, sheet_name="Assigned Inventory", index=False)
            review_queue[OUTPUT_COLUMNS].to_excel(writer, sheet_name="Review Queue", index=False)
            theme_summary[SUMMARY_COLUMNS].to_excel(writer, sheet_name="Theme Summary", index=False)
            self.theme_library_for_display().to_excel(writer, sheet_name="Theme Library", index=False)

            for worksheet in writer.book.worksheets:
                worksheet.freeze_panes = "A2"
                for cells in worksheet.columns:
                    max_length = max(
                        len(str(cell.value)) if cell.value is not None else 0 for cell in cells
                    )
                    worksheet.column_dimensions[cells[0].column_letter].width = min(
                        max(max_length + 2, 12),
                        48,
                    )
        return output.getvalue()

    def build_summary_excel_export(self, theme_summary: pd.DataFrame) -> bytes:
        output = BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            theme_summary[SUMMARY_COLUMNS].to_excel(writer, sheet_name="Theme Summary", index=False)
            worksheet = writer.book["Theme Summary"]
            worksheet.freeze_panes = "A2"
            for cells in worksheet.columns:
                max_length = max(
                    len(str(cell.value)) if cell.value is not None else 0 for cell in cells
                )
                worksheet.column_dimensions[cells[0].column_letter].width = min(
                    max(max_length + 2, 12),
                    56,
                )
        return output.getvalue()

    def enrich_with_llm(
        self,
        assigned: pd.DataFrame,
        theme_summary: pd.DataFrame,
        api_key: str | None = None,
        model: str = "gpt-5.2",
        scope: str = "Top 5 assigned themes",
    ) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
        """Generate optional LLM narratives while preserving deterministic assignments."""
        if assigned.empty or theme_summary.empty:
            raise ValueError("Run theme assignment before generating LLM narratives.")

        resolved_key = api_key or os.getenv("OPENAI_API_KEY", "")
        if not resolved_key:
            raise ValueError("Set OPENAI_API_KEY or enter an API key to use LLM enrichment.")

        try:
            from openai import OpenAI
        except ImportError as exc:
            raise RuntimeError(
                "LLM enrichment requires the OpenAI Python SDK. Install requirements.txt and try again."
            ) from exc

        client = OpenAI(api_key=resolved_key)
        enriched_assigned = assigned.copy()
        enriched_summary = theme_summary.copy()
        for column in ["llm_assignment_justification", "llm_evidence"]:
            if column not in enriched_assigned.columns:
                enriched_assigned[column] = ""
        for column in ["llm_summary", "llm_key_drivers", "llm_management_takeaway"]:
            if column not in enriched_summary.columns:
                enriched_summary[column] = ""

        selected_summary = enriched_summary.copy()
        if "top 5" in scope.lower():
            selected_summary = selected_summary.head(5)

        logs = []
        for _, summary_row in selected_summary.iterrows():
            theme_name = _clean_text(summary_row["theme_name"])
            group = enriched_assigned[
                enriched_assigned.get("final_theme", enriched_assigned["primary_theme"]) == theme_name
            ].copy()
            if group.empty:
                group = enriched_assigned[enriched_assigned["primary_theme"] == theme_name].copy()

            payload = self._build_llm_enrichment_payload(summary_row, group)
            result = self._call_llm_enrichment(client, model, payload)

            summary_mask = enriched_summary["theme_name"] == theme_name
            enriched_summary.loc[summary_mask, "llm_summary"] = result.get("theme_summary", "")
            enriched_summary.loc[summary_mask, "llm_key_drivers"] = "; ".join(
                _as_list(result.get("key_drivers", []))
            )
            enriched_summary.loc[summary_mask, "llm_management_takeaway"] = result.get(
                "management_takeaway",
                "",
            )

            for item in result.get("risk_justifications", []):
                risk_id = _clean_text(item.get("risk_id", ""))
                if not risk_id:
                    continue
                risk_mask = enriched_assigned["risk_id"] == risk_id
                enriched_assigned.loc[risk_mask, "llm_assignment_justification"] = item.get(
                    "justification",
                    "",
                )
                enriched_assigned.loc[risk_mask, "llm_evidence"] = "; ".join(
                    _as_list(item.get("evidence", []))
                )

            logs.append(
                {
                    "theme_name": theme_name,
                    "risks_sent": len(group),
                    "model": model,
                    "status": "Completed",
                }
            )

        return enriched_assigned, enriched_summary, pd.DataFrame(logs)

    def _build_llm_enrichment_payload(
        self,
        summary_row: pd.Series,
        group: pd.DataFrame,
    ) -> dict[str, Any]:
        risk_records = []
        for _, row in group.iterrows():
            risk_records.append(
                {
                    "risk_id": _clean_text(row.get("risk_id", "")),
                    "risk_title": _clean_text(row.get("risk_title", "")),
                    "risk_description": _clean_text(row.get("risk_description", "")),
                    "business_unit": _clean_text(row.get("business_unit", "")),
                    "risk_type": _clean_text(row.get("risk_type", "")),
                    "taxonomy": _clean_text(row.get("taxonomy", "")),
                    "root_cause": _clean_text(row.get("root_cause", "")),
                    "primary_theme": _clean_text(row.get("primary_theme", "")),
                    "primary_theme_score": row.get("primary_theme_score", ""),
                    "secondary_theme": _clean_text(row.get("secondary_theme", "")),
                    "review_flag": _clean_text(row.get("review_flag", "")),
                    "review_reason": _clean_text(row.get("review_reason", "")),
                    "deterministic_rationale": _clean_text(row.get("assignment_rationale", "")),
                }
            )

        return {
            "theme": {
                "theme_name": _clean_text(summary_row.get("theme_name", "")),
                "domain": _clean_text(summary_row.get("domain", "")),
                "total_risks": int(summary_row.get("total_risks", 0)),
                "material_risks": int(summary_row.get("material_risks", 0)),
                "business_units": _clean_text(summary_row.get("business_units", "")),
                "risk_types": _clean_text(summary_row.get("risk_types", "")),
                "top_root_causes": _clean_text(summary_row.get("top_root_causes", "")),
                "rule_based_summary": _clean_text(summary_row.get("ai_summary", "")),
            },
            "risks": risk_records,
        }

    def _call_llm_enrichment(self, client: Any, model: str, payload: dict[str, Any]) -> dict[str, Any]:
        instructions = (
            "You write concise executive risk reporting narratives. Use only the provided data. "
            "Do not invent facts, counts, causes, ratings, or actions. Keep language suitable for "
            "CRO, executive committee, and board reporting. Explain why each risk belongs in the "
            "assigned executive risk theme using evidence from title, description, business unit, "
            "risk type, taxonomy, root cause, score, and candidate themes. Do not expose hidden "
            "chain-of-thought; provide concise justifications and cited evidence snippets."
        )
        input_text = json.dumps(payload, ensure_ascii=False, default=str)
        schema = self._llm_enrichment_schema()

        try:
            response = client.responses.create(
                model=model,
                instructions=instructions,
                input=input_text,
                text={
                    "format": {
                        "type": "json_schema",
                        "name": "risk_theme_enrichment",
                        "schema": schema,
                        "strict": True,
                    }
                },
                max_output_tokens=1800,
                store=False,
            )
        except Exception:
            response = client.responses.create(
                model=model,
                instructions=(
                    instructions
                    + " Return only valid JSON matching this schema: "
                    + json.dumps(schema, ensure_ascii=False)
                ),
                input=input_text,
                max_output_tokens=1800,
                store=False,
            )

        text = self._extract_response_text(response)
        try:
            return json.loads(text)
        except json.JSONDecodeError as exc:
            raise RuntimeError(f"LLM response was not valid JSON: {text[:500]}") from exc

    def _extract_response_text(self, response: Any) -> str:
        output_text = getattr(response, "output_text", None)
        if output_text:
            return output_text

        chunks = []
        for item in getattr(response, "output", []) or []:
            for content in getattr(item, "content", []) or []:
                text = getattr(content, "text", None)
                if text:
                    chunks.append(text)
        if chunks:
            return "\n".join(chunks)
        raise RuntimeError("OpenAI response did not include output text.")

    def _llm_enrichment_schema(self) -> dict[str, Any]:
        return {
            "type": "object",
            "additionalProperties": False,
            "properties": {
                "theme_summary": {
                    "type": "string",
                    "description": "Two to three sentence executive summary of the risk theme group.",
                },
                "key_drivers": {
                    "type": "array",
                    "items": {"type": "string"},
                    "description": "Three to five concise drivers evidenced by the provided risks.",
                },
                "management_takeaway": {
                    "type": "string",
                    "description": "One concise management attention point.",
                },
                "risk_justifications": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "additionalProperties": False,
                        "properties": {
                            "risk_id": {"type": "string"},
                            "justification": {
                                "type": "string",
                                "description": "One to two sentence assignment justification.",
                            },
                            "evidence": {
                                "type": "array",
                                "items": {"type": "string"},
                                "description": "Short evidence phrases from the supplied risk record.",
                            },
                        },
                        "required": ["risk_id", "justification", "evidence"],
                    },
                },
            },
            "required": [
                "theme_summary",
                "key_drivers",
                "management_takeaway",
                "risk_justifications",
            ],
        }

    def build_slide_export(
        self,
        assigned: pd.DataFrame,
        theme_summary: pd.DataFrame,
        scope: str = "Top 5 assigned themes",
        deck_title: str = "AI Risk Theme Assignment Summary",
        deck_subtitle: str | None = None,
        style_preset: str = "Executive Warm",
    ) -> bytes:
        if assigned.empty or theme_summary.empty:
            raise ValueError("Run theme assignment before exporting slides.")

        try:
            from pptx import Presentation
            from pptx.dml.color import RGBColor
            from pptx.enum.shapes import MSO_SHAPE
            from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
            from pptx.util import Inches, Pt
        except ImportError as exc:
            raise RuntimeError(
                "Slides export requires python-pptx. Install requirements.txt and try again."
            ) from exc

        palettes = {
            "Executive Warm": {
                "bg": (245, 242, 236),
                "card": (255, 255, 255),
                "ink": (26, 24, 20),
                "soft": (92, 86, 80),
                "line": (227, 221, 210),
                "line_soft": (236, 235, 230),
                "accent": (214, 69, 69),
                "green": (47, 125, 109),
            },
            "Board Minimal": {
                "bg": (248, 248, 246),
                "card": (255, 255, 255),
                "ink": (22, 24, 27),
                "soft": (88, 93, 100),
                "line": (214, 218, 224),
                "line_soft": (235, 237, 240),
                "accent": (172, 47, 47),
                "green": (43, 111, 93),
            },
            "Slate Contrast": {
                "bg": (239, 241, 238),
                "card": (255, 255, 255),
                "ink": (18, 25, 31),
                "soft": (82, 92, 99),
                "line": (205, 213, 214),
                "line_soft": (230, 235, 233),
                "accent": (190, 56, 46),
                "green": (31, 116, 109),
            },
        }
        palette = palettes.get(style_preset, palettes["Executive Warm"])
        bg = RGBColor(*palette["bg"])
        card = RGBColor(*palette["card"])
        ink = RGBColor(*palette["ink"])
        soft = RGBColor(*palette["soft"])
        line = RGBColor(*palette["line"])
        line_soft = RGBColor(*palette["line_soft"])
        accent = RGBColor(*palette["accent"])
        green = RGBColor(*palette["green"])

        def clean(value: Any, limit: int | None = None) -> str:
            text = _clean_text(value)
            if limit and len(text) > limit:
                return text[: max(limit - 1, 0)].rstrip() + "..."
            return text

        def set_background(slide) -> None:
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = bg

        def add_text(
            slide,
            text: str,
            x: float,
            y: float,
            w: float,
            h: float,
            size: float = 14,
            bold: bool = False,
            color=ink,
            align=PP_ALIGN.LEFT,
        ):
            shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            frame = shape.text_frame
            frame.clear()
            frame.margin_left = Inches(0.02)
            frame.margin_right = Inches(0.02)
            frame.margin_top = Inches(0.02)
            frame.margin_bottom = Inches(0.02)
            frame.word_wrap = True
            frame.vertical_anchor = MSO_ANCHOR.TOP
            paragraph = frame.paragraphs[0]
            paragraph.alignment = align
            run = paragraph.add_run()
            run.text = text
            run.font.name = "Aptos"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
            return shape

        def add_card(slide, x: float, y: float, w: float, h: float):
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x),
                Inches(y),
                Inches(w),
                Inches(h),
            )
            shape.fill.solid()
            shape.fill.fore_color.rgb = card
            shape.line.color.rgb = line
            shape.line.width = Pt(0.8)
            return shape

        def set_cell(cell, text: str, size: float, bold: bool = False, fill=None, font=ink):
            if fill is not None:
                cell.fill.solid()
                cell.fill.fore_color.rgb = fill
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            frame = cell.text_frame
            frame.clear()
            frame.margin_left = Inches(0.04)
            frame.margin_right = Inches(0.04)
            frame.margin_top = Inches(0.02)
            frame.margin_bottom = Inches(0.02)
            paragraph = frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.LEFT
            run = paragraph.add_run()
            run.text = text
            run.font.name = "Aptos"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = font

        def add_table(
            slide,
            rows: list[list[str]],
            x: float,
            y: float,
            widths: list[float],
            height: float,
            font_size: float = 8.0,
        ):
            table_shape = slide.shapes.add_table(
                len(rows),
                len(widths),
                Inches(x),
                Inches(y),
                Inches(sum(widths)),
                Inches(height),
            )
            table = table_shape.table
            for index, width in enumerate(widths):
                table.columns[index].width = Inches(width)

            row_height = max(0.12, height / max(len(rows), 1))
            for row_index, row_values in enumerate(rows):
                table.rows[row_index].height = Inches(row_height)
                for col_index, value in enumerate(row_values):
                    is_header = row_index == 0
                    set_cell(
                        table.cell(row_index, col_index),
                        value,
                        size=max(5.5, font_size - (0 if is_header else 0.4)),
                        bold=is_header,
                        fill=line_soft if is_header else card,
                        font=ink if is_header else soft,
                    )
            return table_shape

        selected_summary = theme_summary.copy()
        if "top 5" in scope.lower():
            selected_summary = selected_summary.head(5)

        working = assigned.copy()
        if "final_theme" not in working.columns:
            working["final_theme"] = working["primary_theme"]

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank_layout = prs.slide_layouts[6]

        total_risks = len(working)
        review_items = int((working["review_flag"] == "Yes").sum())
        auto_items = int((working["review_flag"] == "No").sum())
        material_risks = int(working.get("materiality", pd.Series(dtype=str)).apply(self._is_material).sum())
        themes_used = int(working["final_theme"].nunique())

        slide = prs.slides.add_slide(blank_layout)
        set_background(slide)
        subtitle = deck_subtitle or f"{scope} | Vector retrieval plus metadata scoring"
        add_text(slide, clean(deck_title, 80), 0.55, 0.34, 8.4, 0.45, 24, True)
        add_text(
            slide,
            clean(subtitle, 120),
            0.58,
            0.83,
            7.5,
            0.25,
            9.5,
            False,
            soft,
        )
        add_text(slide, "RT", 12.25, 0.35, 0.55, 0.3, 12, True, accent, PP_ALIGN.RIGHT)

        metric_data = [
            ("Total Risks", f"{total_risks:,}"),
            ("Themes Used", f"{themes_used:,}"),
            ("Auto Assigned", f"{auto_items:,}"),
            ("Review Items", f"{review_items:,}"),
            ("Material Risks", f"{material_risks:,}"),
        ]
        for index, (label, value) in enumerate(metric_data):
            x = 0.55 + index * 2.48
            add_card(slide, x, 1.28, 2.18, 0.78)
            add_text(slide, label, x + 0.16, 1.43, 1.75, 0.18, 7.8, True, soft)
            add_text(slide, value, x + 0.16, 1.62, 1.75, 0.25, 17, True, ink)

        add_text(slide, "Theme Portfolio", 0.55, 2.38, 3.0, 0.28, 14, True)
        top_rows = [["Theme", "Domain", "Risks", "Material", "Root Causes"]]
        for _, row in selected_summary.head(8).iterrows():
            top_rows.append(
                [
                    clean(row["theme_name"], 34),
                    clean(row["domain"], 24),
                    str(int(row["total_risks"])),
                    str(int(row["material_risks"])),
                    clean(row["top_root_causes"], 55),
                ]
            )
        add_table(slide, top_rows, 0.55, 2.78, [3.2, 2.1, 0.7, 0.85, 4.6], 2.35, 8.0)

        add_card(slide, 0.55, 5.5, 12.15, 1.05)
        add_text(slide, "Executive Readout", 0.78, 5.72, 2.3, 0.25, 9, True, accent)
        readout = (
            f"The assignment engine grouped {total_risks:,} risks into {themes_used:,} executive themes. "
            f"{auto_items:,} risks met the auto-assignment rule, while {review_items:,} require human review. "
            f"The following slides show the selected themes and all underlying assigned risks."
        )
        add_text(slide, readout, 0.78, 5.98, 11.5, 0.35, 10.5, False, ink)

        for _, summary_row in selected_summary.iterrows():
            theme_name = clean(summary_row["theme_name"])
            group = working[working["final_theme"] == theme_name].copy()
            if group.empty:
                group = working[working["primary_theme"] == theme_name].copy()
            group = group.sort_values(["review_flag", "primary_theme_score"], ascending=[False, False])

            slide = prs.slides.add_slide(blank_layout)
            set_background(slide)
            add_text(slide, theme_name, 0.55, 0.34, 8.9, 0.42, 22, True)
            add_text(
                slide,
                f"{clean(summary_row['domain'])} | {int(summary_row['total_risks'])} risks | "
                f"{int(summary_row['material_risks'])} material | Review items: "
                f"{int((group['review_flag'] == 'Yes').sum())}",
                0.58,
                0.82,
                8.9,
                0.24,
                9.2,
                False,
                soft,
            )
            add_text(slide, "RT", 12.25, 0.35, 0.55, 0.3, 12, True, accent, PP_ALIGN.RIGHT)

            add_card(slide, 0.55, 1.26, 7.25, 1.28)
            add_text(slide, "AI Summary", 0.77, 1.47, 1.4, 0.18, 8.2, True, accent)
            summary_text = summary_row.get("llm_summary", "") or summary_row["ai_summary"]
            add_text(slide, clean(summary_text, 310), 0.77, 1.72, 6.75, 0.55, 10.0, False, ink)

            add_card(slide, 8.05, 1.26, 4.65, 1.28)
            add_text(slide, "Signals", 8.25, 1.47, 1.2, 0.18, 8.2, True, accent)
            driver_text = summary_row.get("llm_key_drivers", "") or summary_row["top_root_causes"]
            signals = (
                f"Business units: {clean(summary_row['business_units'], 86)}\n"
                f"Risk types: {clean(summary_row['risk_types'], 86)}\n"
                f"Drivers: {clean(driver_text, 86)}"
            )
            add_text(slide, signals, 8.25, 1.72, 4.1, 0.55, 8.4, False, soft)

            add_text(slide, "Underlying Risks", 0.55, 2.87, 2.3, 0.25, 13.5, True)
            risk_rows = [["ID", "Risk Title", "Business", "Type", "Score", "Review"]]
            title_limit = 92 if len(group) <= 12 else 68
            for _, risk_row in group.iterrows():
                risk_rows.append(
                    [
                        clean(risk_row["risk_id"], 12),
                        clean(risk_row["risk_title"], title_limit),
                        clean(risk_row["business_unit"], 22),
                        clean(risk_row["risk_type"], 16),
                        f"{float(risk_row['primary_theme_score']):.2f}",
                        clean(risk_row["review_flag"], 5),
                    ]
                )
            table_font = 8.0 if len(risk_rows) <= 12 else max(5.7, 9.0 - len(risk_rows) * 0.16)
            add_table(
                slide,
                risk_rows,
                0.55,
                3.22,
                [0.78, 5.1, 1.72, 1.18, 0.75, 0.75],
                3.72,
                table_font,
            )

        output = BytesIO()
        prs.save(output)
        return output.getvalue()

    def _theme_metadata_from_library(self, themes: pd.DataFrame) -> list[dict[str, Any]]:
        metadata = []
        for _, row in themes.iterrows():
            theme = row.to_dict()
            item = {
                "theme_id": theme["theme_id"],
                "theme_name": theme["theme_name"],
                "domain": theme["domain"],
                "description": theme["description"],
                "typical_business_units": _as_list(theme.get("typical_business_units", [])),
                "typical_risk_types": _as_list(theme.get("typical_risk_types", [])),
                "root_causes": _as_list(theme.get("root_causes", [])),
                "keywords": _as_list(theme.get("keywords", [])),
            }
            item["theme_document"] = self.build_theme_document(item)
            metadata.append(item)
        return metadata

    def _load_faiss_vector_store(self) -> bool:
        try:
            self.embedding_model = self._load_embedding_model()
            if self.embedding_model is None:
                return False

            import faiss

            self.faiss = faiss

            with self.metadata_path.open("r", encoding="utf-8") as file:
                self.theme_metadata = json.load(file)
            self.theme_documents = [item["theme_document"] for item in self.theme_metadata]
            self.faiss_index = faiss.read_index(str(self.index_path))
            self.backend = f"FAISS + sentence-transformers/{self.embedding_model_name}"
            self.status_message = "Loaded existing FAISS theme vector store."
            self.ready = True
            return True
        except Exception as exc:
            self.status_message = f"FAISS load unavailable, using TF-IDF fallback: {exc}"
            self.faiss_index = None
            self.embedding_model = None
            return False

    def _build_faiss_vector_store(self) -> bool:
        try:
            self.embedding_model = self._load_embedding_model()
            if self.embedding_model is None:
                return False

            vectors = self.embedding_model.encode(
                self.theme_documents,
                normalize_embeddings=True,
                show_progress_bar=False,
            ).astype("float32")

            import faiss

            self.faiss = faiss
            index = faiss.IndexFlatIP(vectors.shape[1])
            index.add(vectors)

            self.vector_store_dir.mkdir(parents=True, exist_ok=True)
            faiss.write_index(index, str(self.index_path))
            with self.metadata_path.open("w", encoding="utf-8") as file:
                json.dump(self.theme_metadata, file, indent=2)

            self.faiss_index = index
            self.backend = f"FAISS + sentence-transformers/{self.embedding_model_name}"
            self.status_message = "Built FAISS theme vector store from risk_theme_library.yaml."
            self.ready = True
            return True
        except Exception as exc:
            self.status_message = f"FAISS build unavailable, using TF-IDF fallback: {exc}"
            self.faiss_index = None
            self.embedding_model = None
            return False

    def _load_embedding_model(self):
        try:
            from sentence_transformers import SentenceTransformer

            try:
                return SentenceTransformer(self.embedding_model_name, local_files_only=True)
            except TypeError:
                return SentenceTransformer(self.embedding_model_name)
            except Exception:
                return SentenceTransformer(self.embedding_model_name)
        except Exception as exc:
            self.status_message = (
                "sentence-transformers model unavailable, using TF-IDF fallback: "
                f"{str(exc).splitlines()[0]}"
            )
            return None

    def _build_tfidf_fallback(self) -> None:
        self.vectorizer = TfidfVectorizer(
            stop_words="english",
            ngram_range=(1, 2),
            min_df=1,
            sublinear_tf=True,
        )
        self.theme_tfidf_matrix = self.vectorizer.fit_transform(self.theme_documents)
        self.backend = "TF-IDF cosine similarity fallback"
        self.ready = True
        if not self.status_message:
            self.status_message = "Using TF-IDF fallback because FAISS or embeddings are unavailable."

    def _score_candidate(
        self,
        risk: pd.Series,
        risk_profile: str,
        candidate: dict[str, Any],
    ) -> dict[str, Any]:
        semantic = float(candidate.get("semantic_similarity", 0.0))
        business = self._field_match_score(
            risk.get("business_unit", ""),
            candidate.get("typical_business_units", []),
        )
        risk_type = self._field_match_score(
            risk.get("risk_type", ""),
            candidate.get("typical_risk_types", []),
        )
        root_cause, matched_root_causes = self._root_cause_score(
            risk,
            risk_profile,
            candidate.get("root_causes", []),
        )
        keyword, matched_keywords = self._keyword_score(
            risk_profile,
            candidate.get("keywords", []),
        )

        final_score = (
            0.50 * semantic
            + 0.20 * business
            + 0.15 * risk_type
            + 0.10 * root_cause
            + 0.05 * keyword
        )

        scored = dict(candidate)
        scored.update(
            {
                "semantic_similarity": semantic,
                "business_match_score": business,
                "risk_type_match_score": risk_type,
                "root_cause_match_score": root_cause,
                "keyword_match_score": keyword,
                "final_score": float(final_score),
                "matched_root_causes": matched_root_causes,
                "matched_keywords": matched_keywords,
            }
        )
        return scored

    def _field_match_score(self, value: Any, candidates: list[str]) -> float:
        value_norm = _normalize_text(value)
        if not value_norm or not candidates:
            return 0.0
        value_tokens = _tokens(value_norm)

        best = 0.0
        for candidate in candidates:
            candidate_norm = _normalize_text(candidate)
            candidate_tokens = _tokens(candidate_norm)
            if not candidate_norm or not candidate_tokens:
                continue
            if value_norm == candidate_norm:
                best = max(best, 1.0)
            elif candidate_norm in value_norm or value_norm in candidate_norm:
                best = max(best, 0.9)
            else:
                overlap = len(value_tokens & candidate_tokens)
                if overlap:
                    best = max(best, min(0.8, overlap / max(len(candidate_tokens), 1)))
        return float(best)

    def _root_cause_score(
        self,
        risk: pd.Series,
        risk_profile: str,
        theme_root_causes: list[str],
    ) -> tuple[float, list[str]]:
        root_text = _normalize_text(risk.get("root_cause", ""))
        taxonomy_text = _normalize_text(risk.get("taxonomy", ""))
        profile_text = _normalize_text(risk_profile)
        combined = " ".join([root_text, taxonomy_text, profile_text])
        matched = []
        best = 0.0

        for candidate in theme_root_causes:
            candidate_norm = _normalize_text(candidate)
            if not candidate_norm:
                continue
            if root_text and root_text == candidate_norm:
                best = max(best, 1.0)
                matched.append(candidate)
            elif candidate_norm in combined:
                best = max(best, 0.9)
                matched.append(candidate)
            else:
                overlap = len(_tokens(root_text + " " + taxonomy_text) & _tokens(candidate_norm))
                if overlap:
                    best = max(best, min(0.8, overlap / max(len(_tokens(candidate_norm)), 1)))
                    matched.append(candidate)

        return float(best), matched[:4]

    def _keyword_score(self, risk_profile: str, keywords: list[str]) -> tuple[float, list[str]]:
        profile_norm = _normalize_text(risk_profile)
        profile_tokens = _tokens(profile_norm)
        matched = []
        for keyword in keywords:
            keyword_norm = _normalize_text(keyword)
            keyword_tokens = _tokens(keyword_norm)
            if not keyword_norm:
                continue
            if keyword_norm in profile_norm or keyword_tokens <= profile_tokens:
                matched.append(keyword)

        denominator = max(min(len(keywords), 4), 1)
        return float(min(1.0, len(matched) / denominator)), matched[:6]

    def _assignment_rationale(self, candidate: dict[str, Any]) -> str:
        signals = [f"semantic similarity {candidate['semantic_similarity']:.2f}"]
        if candidate["business_match_score"] >= 0.8:
            signals.append("business unit alignment")
        if candidate["risk_type_match_score"] >= 0.8:
            signals.append("risk type alignment")
        if candidate["root_cause_match_score"] >= 0.8:
            signals.append("taxonomy/root cause alignment")
        if candidate.get("matched_keywords"):
            signals.append("keywords: " + ", ".join(candidate["matched_keywords"]))
        return f"Assigned to {candidate['theme_name']} based on " + "; ".join(signals) + "."

    def _rule_based_summary(
        self,
        theme_name: str,
        group: pd.DataFrame,
        business_units: str,
        risk_types: str,
        root_causes: str,
        material_risks: int,
    ) -> str:
        total = len(group)
        return (
            f"{theme_name} includes {total} risk{'s' if total != 1 else ''} across "
            f"{business_units or 'the represented business units'}. Key drivers include "
            f"{root_causes or 'the listed root causes'}. The theme spans "
            f"{risk_types or 'multiple risk types'}, with {material_risks} material "
            f"risk{'s' if material_risks != 1 else ''} requiring management attention."
        )

    def _join_unique(self, values: pd.Series) -> str:
        unique = sorted({_clean_text(value) for value in values if _clean_text(value)})
        return ", ".join(unique)

    def _top_values(self, values: pd.Series, limit: int = 4) -> str:
        counter = Counter(_clean_text(value) for value in values if _clean_text(value))
        return ", ".join(value for value, _ in counter.most_common(limit))

    def _is_material(self, value: Any) -> bool:
        normalized = _normalize_text(value)
        if not normalized or normalized.startswith("non"):
            return False
        return normalized in {"material", "yes", "y", "true", "1", "high"}
