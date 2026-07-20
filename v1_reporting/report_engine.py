from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Iterable

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


@dataclass
class ReportPlan:
    inventory: str = "Q2 2026"
    comparison: str = "Q1 2026"
    legal_entity: str = "CUSO"
    business_division: str = "Investment Bank"
    risk_population: str = "Material risks only"
    audience: str = "Executive Leadership"
    sections: list[str] = field(default_factory=list)


DEFAULT_SECTIONS = [
    "Executive Summary",
    "Portfolio Overview",
    "Top Risks",
    "QoQ Changes",
    "Management Actions",
]


BUSINESS_ALIASES = {
    "investment bank": "Investment Bank",
    "ib": "Investment Bank",
    "wealth management": "Wealth Management",
    "wm": "Wealth Management",
    "asset management": "Asset Management",
    "corporate lending": "Corporate Lending",
    "global markets": "Global Markets",
    "treasury": "Treasury",
}


def _contains_token(text: str, token: str) -> bool:
    return bool(re.search(rf"(?<!\w){re.escape(token)}(?!\w)", text, flags=re.IGNORECASE))


def infer_plan_from_goal(goal: str) -> ReportPlan:
    """Deterministic MVP parser.

    Replace this function with an LLM structured-output call later. The rest of
    the UI and workflow can stay unchanged because both approaches return the
    same ReportPlan object.
    """
    text = goal.lower().strip()
    plan = ReportPlan(sections=DEFAULT_SECTIONS.copy())

    for alias, business in BUSINESS_ALIASES.items():
        if _contains_token(text, alias):
            plan.business_division = business
            break

    if plan.business_division in {"Corporate Lending", "Global Markets", "Treasury"}:
        plan.legal_entity = "US Branches"

    if "ah llc" in text:
        plan.legal_entity = "AH LLC"
    elif "group" in text:
        plan.legal_entity = "Group"
    elif "cuso" in text:
        plan.legal_entity = "CUSO"

    if "all business" in text or "enterprise" in text:
        plan.business_division = "All Businesses"

    if any(term in text for term in ["new risk", "changed risk", "material change", "changes"]):
        plan.risk_population = "New and changed risks"

    if "non-material" in text or "non material" in text:
        plan.risk_population = "Non-material risks only"
    elif "all active" in text:
        plan.risk_population = "All active risks"

    if "regulator" in text or "regulatory" in text:
        plan.audience = "Regulator"
        _append_unique(plan.sections, "Risk ID Appendix")
    elif "cro" in text:
        plan.audience = "CRO"
    elif "risk manager" in text:
        plan.audience = "Risk Manager"

    if any(term in text for term in ["risk id", "appendix", "audit"]):
        _append_unique(plan.sections, "Risk ID Appendix")
    if any(term in text for term in ["quantification", "exposure"]):
        _append_unique(plan.sections, "Risk Quantification")
    if "concentration" in text:
        _append_unique(plan.sections, "Risk Concentrations")

    return plan


def _append_unique(items: list[str], value: str) -> None:
    if value not in items:
        items.append(value)


def build_refined_goal(plan: ReportPlan) -> str:
    comparison_sentence = (
        "Focus on the current inventory without a period comparison."
        if plan.comparison == "No comparison"
        else f"Compare {plan.inventory} with {plan.comparison}."
    )
    return (
        f"Generate a {plan.audience.lower()} risk report for {plan.legal_entity}, "
        f"covering {plan.business_division}, using the {plan.inventory} inventory. "
        f"{comparison_sentence} Focus on {plan.risk_population.lower()}, highlight "
        "the most significant risks and changes, and provide concise management actions."
    )


def build_default_filename(plan: ReportPlan) -> str:
    abbreviations = {
        "Investment Bank": "IB",
        "Wealth Management": "WM",
        "Asset Management": "AM",
        "Corporate Lending": "CL",
        "Global Markets": "GM",
        "Corporate Center": "CC",
        "All Businesses": "All",
    }
    business = abbreviations.get(plan.business_division, plan.business_division)
    raw = f"Risk_Atlas_{plan.legal_entity}_{business}_{plan.inventory}.pptx"
    return sanitize_filename(raw)


def sanitize_filename(filename: str) -> str:
    base = re.sub(r"\.pptx$", "", filename.strip(), flags=re.IGNORECASE)
    base = re.sub(r"[^A-Za-z0-9._-]+", "_", base)
    base = re.sub(r"_+", "_", base).strip("._-")
    if not base:
        base = "Risk_Atlas_Report"
    return f"{base}.pptx"


def generate_powerpoint(
    plan: ReportPlan,
    refined_goal: str,
    sections: Iterable[str],
) -> bytes:
    """Create a compact, real PPTX package for the demo.

    Production implementation can replace placeholder bullets with inventory
    query results, Plotly chart images and source-risk tables.
    """
    sections = list(sections)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Cover
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = f"{plan.business_division} Risk Report"
    slide.placeholders[1].text = (
        f"{plan.legal_entity} | {plan.inventory} | {plan.audience}\n"
        f"Generated by Risk Atlas"
    )
    _format_slide(slide)

    # Confirmed scope
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Confirmed Report Scope"
    scope_lines = [
        f"Legal entity: {plan.legal_entity}",
        f"Business division: {plan.business_division}",
        f"Inventory: {plan.inventory}",
        f"Comparison: {plan.comparison}",
        f"Risk population: {plan.risk_population}",
        f"Audience: {plan.audience}",
        f"Report goal: {refined_goal}",
    ]
    _set_bullets(slide.placeholders[1].text_frame, scope_lines)
    _format_slide(slide)

    # One slide per selected section
    for section in sections:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = section
        bullets = _placeholder_content(section, plan)
        _set_bullets(slide.placeholders[1].text_frame, bullets)
        _format_slide(slide)

    # Audit footer slide
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Review and Audit Notes"
    _set_bullets(
        slide.placeholders[1].text_frame,
        [
            "Scope was confirmed before generation.",
            "Selected report sections were retained in the generation request.",
            "Production version should attach source Risk IDs to every material statement.",
            "Human review is required before external or regulatory distribution.",
        ],
    )
    _format_slide(slide)

    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output.read()


def _placeholder_content(section: str, plan: ReportPlan) -> list[str]:
    content = {
        "Executive Summary": [
            f"Executive view of {plan.business_division} within {plan.legal_entity}.",
            "Highlight material movements, concentrations and decisions required.",
            "Replace these placeholders with generated, source-grounded findings.",
        ],
        "Portfolio Overview": [
            f"Population: {plan.risk_population}.",
            f"Current inventory: {plan.inventory}.",
            f"Comparison inventory: {plan.comparison}.",
        ],
        "Top Risks": [
            "Rank risks using a deterministic business-approved metric.",
            "Show rank, exposure, rating, materiality and source Risk ID.",
            "Explain why each item is included in the top set.",
        ],
        "QoQ Changes": [
            "Identify new, retired, upgraded and downgraded risks.",
            "Separate data changes from genuine risk-profile changes.",
            "Show prior and current values side by side.",
        ],
        "Risk Quantification": [
            "Compare exposure, limits, stress metrics and management thresholds.",
            "Flag missing or inconsistent quantification fields.",
            "Retain calculation methodology for audit review.",
        ],
        "Risk Concentrations": [
            "Highlight concentrations by taxonomy, business and legal entity.",
            "Distinguish material concentration from high record count.",
            "Provide management interpretation and trend context.",
        ],
        "Management Actions": [
            "Summarize decisions, owners and target completion dates.",
            "Separate confirmed actions from AI suggestions.",
            "Escalate unresolved data gaps or review items.",
        ],
        "Risk ID Appendix": [
            "List source Risk IDs and inventory attributes used in the report.",
            "Provide direct links or downloadable inventory extracts in production.",
            "Preserve the confirmed scope and generation timestamp.",
        ],
    }
    return content.get(section, ["Section content will be generated from the confirmed inventory scope."])


def _set_bullets(text_frame, lines: list[str]) -> None:
    text_frame.clear()
    for idx, line in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if idx == 0 else text_frame.add_paragraph()
        paragraph.text = line
        paragraph.level = 0
        paragraph.font.size = Pt(20)
        paragraph.space_after = Pt(8)


def _format_slide(slide) -> None:
    title = slide.shapes.title
    if title is not None:
        title.text_frame.paragraphs[0].font.name = "Arial"
        title.text_frame.paragraphs[0].font.size = Pt(28)
        title.text_frame.paragraphs[0].font.bold = True
        title.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
